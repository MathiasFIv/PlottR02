#!/usr/bin/env python3
"""Generate a minimal load-test fixture workbook.

The API's ClosedXmlRatingRepository reads the sheet named "Subsample Plot Data"
with a fixed layout (see ClosedXmlRatingRepository.TraitColumn = ordinal + 2):
  * trait ordinals 1..9 -> columns C..K (i.e. 3..11)
  * row 3  -> pest codes (columns C..K, i.e. 3..11)
  * row 5  -> crop
  * row 8  -> Part Rated category labels
  * row 20 -> subsample count
  * data rows start at row 22; column B holds "plot subsample" keys

This builds a workbook with the required sheet name and a set of empty data
rows (plots 101..114, subsamples 1..4). Empty cells matter: the smoke/steady
writes use `overwrite: true` so writes are idempotent regardless, but an empty
sheet keeps the first write of every cell a real write.

It is only a fallback: when the archived fixtures exist the workflow copies one
of those instead. Requires openpyxl (pip install openpyxl).
"""

import sys

from openpyxl import Workbook

SHEET_NAME = 'Subsample Plot Data'
FIRST_DATA_ROW = 22
PEST_ROW = 3
CROP_ROW = 5
PART_RATED_ROW = 8
SUBSAMPLE_COUNT_ROW = 20
FIRST_TRAIT_COLUMN = 3  # column C = trait ordinal 1 (TraitColumn = ordinal + 2)
LAST_TRAIT_COLUMN = 11  # column K = trait ordinal 9

PEST_CODES = ['ERYSGR', 'PUCCST', 'SEPTTR', 'ERYSGR', 'PUCCST', 'SEPTTR', 'ERYSGR', 'PUCCST', 'SEPTTR']
PART_RATED = ['LEAF1', 'LEAF1', 'LEAF2', 'LEAF2', 'LEAF2', 'LEAF3', 'LEAF3', 'LEAF3', 'LEAF3']
PLOTS = list(range(101, 115))  # 101..114
SUBSAMPLES = [1, 2, 3, 4]


def build(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for i, code in enumerate(PEST_CODES):
        ws.cell(PEST_ROW, FIRST_TRAIT_COLUMN + i, code)
    for i, label in enumerate(PART_RATED):
        ws.cell(PART_RATED_ROW, FIRST_TRAIT_COLUMN + i, label)
    for col in range(FIRST_TRAIT_COLUMN, LAST_TRAIT_COLUMN + 1):
        ws.cell(CROP_ROW, col, 'Winter wheat')

    ws.cell(SUBSAMPLE_COUNT_ROW, FIRST_TRAIT_COLUMN, len(SUBSAMPLES))

    row = FIRST_DATA_ROW
    for plot in PLOTS:
        for subsample in SUBSAMPLES:
            ws.cell(row, 2, f'{plot} {subsample}')  # column B key "plot subsample"
            row += 1

    wb.save(path)
    print(f'Wrote fixture with {len(PLOTS) * len(SUBSAMPLES)} data rows to {path}')


if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else 'fixtures/trial.xlsx'
    build(target)
